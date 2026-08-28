"""RAG 文档加载器（多格式 / 多模态）"""
import logging
import csv
import os
import re
import asyncio
import concurrent.futures


def _run_async(coro):
    """在同步上下文中运行协程；若当前已在事件循环内则开线程执行，避免 "loop already running"。"""
    try:
        asyncio.get_running_loop()
    except RuntimeError:
        return asyncio.run(coro)
    with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ex:
        return ex.submit(lambda: asyncio.run(coro)).result()
from html.parser import HTMLParser
from typing import List, Dict, Any
from pathlib import Path
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass
class Document:
    """文档对象"""
    content: str
    metadata: Dict[str, Any]
    doc_id: str = ""


class _HTMLTextExtractor(HTMLParser):
    """极简 HTML 标签剥离器（无外部依赖）"""

    def __init__(self):
        super().__init__()
        self._parts: List[str] = []

    def handle_data(self, data):
        self._parts.append(data)

    def get_text(self) -> str:
        return "\n".join(p.strip() for p in self._parts if p.strip())


class DocumentLoader:
    """文档加载器"""

    # 支持的后缀 -> 处理方法名
    _DISPATCH = {
        ".pdf": "load_pdf",
        ".docx": "load_docx",
        ".txt": "load_txt",
        ".md": "load_md",
        ".markdown": "load_md",
        ".html": "load_html",
        ".htm": "load_html",
        ".csv": "load_csv",
        ".xlsx": "load_xlsx",
        ".png": "load_image",
        ".jpg": "load_image",
        ".jpeg": "load_image",
        ".gif": "load_image",
        ".bmp": "load_image",
        ".webp": "load_image",
        ".pptx": "load_pptx",
    }

    @staticmethod
    def load_pdf(file_path: str) -> List[Document]:
        """加载 PDF 文件"""
        try:
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            documents = []

            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text.strip():
                    documents.append(Document(
                        content=text,
                        metadata={
                            "source": file_path,
                            "page": i + 1,
                            "type": "pdf"
                        }
                    ))

            logger.info(f"Loaded {len(documents)} pages from {file_path}")
            return documents

        except Exception as e:
            logger.error(f"Failed to load PDF {file_path}: {e}")
            return []

    @staticmethod
    def load_docx(file_path: str) -> List[Document]:
        """加载 Word 文档"""
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(file_path)
            paragraphs = []

            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            if paragraphs:
                return [Document(
                    content="\n\n".join(paragraphs),
                    metadata={
                        "source": file_path,
                        "type": "docx"
                    }
                )]

            return []

        except Exception as e:
            logger.error(f"Failed to load DOCX {file_path}: {e}")
            return []

    @staticmethod
    def load_txt(file_path: str) -> List[Document]:
        """加载文本文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            if content.strip():
                return [Document(
                    content=content,
                    metadata={
                        "source": file_path,
                        "type": "txt"
                    }
                )]

            return []

        except Exception as e:
            logger.error(f"Failed to load TXT {file_path}: {e}")
            return []

    @staticmethod
    def load_md(file_path: str) -> List[Document]:
        """加载 Markdown 文件"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            if content.strip():
                return [Document(
                    content=content,
                    metadata={"source": file_path, "type": "markdown"}
                )]
            return []
        except Exception as e:
            logger.error(f"Failed to load MD {file_path}: {e}")
            return []

    @staticmethod
    def load_html(file_path: str) -> List[Document]:
        """加载 HTML 文件（剥离标签）"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                raw = f.read()
            parser = _HTMLTextExtractor()
            parser.feed(raw)
            text = parser.get_text()
            if text.strip():
                return [Document(
                    content=text,
                    metadata={"source": file_path, "type": "html"}
                )]
            return []
        except Exception as e:
            logger.error(f"Failed to load HTML {file_path}: {e}")
            return []

    @staticmethod
    def load_csv(file_path: str) -> List[Document]:
        """加载 CSV：每行作为一条记录文本"""
        try:
            docs: List[Document] = []
            with open(file_path, 'r', encoding='utf-8', errors='ignore', newline='') as f:
                reader = csv.reader(f)
                header = next(reader, None)
                for i, row in enumerate(reader, 1):
                    if not any(c.strip() for c in row):
                        continue
                    if header:
                        line = "; ".join(f"{h}: {c}" for h, c in zip(header, row))
                    else:
                        line = ", ".join(row)
                    docs.append(Document(
                        content=line,
                        metadata={"source": file_path, "type": "csv", "row": i}
                    ))
            logger.info(f"Loaded {len(docs)} rows from CSV {file_path}")
            return docs
        except Exception as e:
            logger.error(f"Failed to load CSV {file_path}: {e}")
            return []

    @staticmethod
    def load_xlsx(file_path: str) -> List[Document]:
        """加载 Excel：每个工作表每行一条记录"""
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.warning("openpyxl 未安装，无法解析 xlsx；请 pip install openpyxl")
            return []

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            docs: List[Document] = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                rows = list(ws.iter_rows(values_only=True))
                header = [str(c) if c is not None else "" for c in rows[0]] if rows else []
                for i, row in enumerate(rows[1:], 1):
                    if not any(c is not None and str(c).strip() for c in row):
                        continue
                    cells = [str(c) if c is not None else "" for c in row]
                    if header:
                        line = "; ".join(f"{h}: {c}" for h, c in zip(header, cells))
                    else:
                        line = ", ".join(cells)
                    docs.append(Document(
                        content=line,
                        metadata={"source": file_path, "type": "xlsx", "sheet": sheet, "row": i}
                    ))
            logger.info(f"Loaded {len(docs)} rows from XLSX {file_path}")
            return docs
        except Exception as e:
            logger.error(f"Failed to load XLSX {file_path}: {e}")
            return []

    @staticmethod
    def load_pptx(file_path: str) -> List[Document]:
        """加载 PPTX：每页作为一条文档"""
        try:
            from pptx import Presentation
        except ImportError:
            logger.warning("python-pptx 未安装，无法解析 pptx；请 pip install python-pptx")
            return []

        try:
            prs = Presentation(file_path)
            docs: List[Document] = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text.strip()]
                if texts:
                    docs.append(Document(
                        content="\n".join(texts),
                        metadata={"source": file_path, "type": "pptx", "slide": i}
                    ))
            logger.info(f"Loaded {len(docs)} slides from PPTX {file_path}")
            return docs
        except Exception as e:
            logger.error(f"Failed to load PPTX {file_path}: {e}")
            return []

    @staticmethod
    def load_image(file_path: str) -> List[Document]:
        """加载图片：OCR 提取文本（复用 core.multimodal.image_processor）"""
        try:
            from core.multimodal.image_processor import image_processor
            with open(file_path, "rb") as f:
                data = f.read()
            result = _run_async(image_processor.process_image(data, extract_text=True))
            text = result.get("extracted_text", "") or ""
            if text.strip():
                return [Document(
                    content=text,
                    metadata={
                        "source": file_path,
                        "type": "image_ocr",
                        "width": result.get("width"),
                        "height": result.get("height"),
                    }
                )]
            logger.warning(f"图片 OCR 无文本: {file_path}（可能未装 pytesseract 或图片无文字）")
            return []
        except Exception as e:
            logger.error(f"Failed to load image {file_path}: {e}")
            return []

    @classmethod
    def load(cls, file_path: str) -> List[Document]:
        """根据文件类型加载"""
        path = Path(file_path)
        suffix = path.suffix.lower()

        handler = cls._DISPATCH.get(suffix)
        if not handler:
            logger.warning(f"Unsupported file type: {suffix}")
            return []

        return getattr(cls, handler)(file_path)
